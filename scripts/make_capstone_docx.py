# -*- coding: utf-8 -*-
"""캡스톤디자인 결과 — 항목별 내용을 워드(.docx)로 작성한다.

사용자가 PPT 템플릿에 직접 옮겨 적을 수 있도록, 4개 핵심 항목
(개발 동기 및 목적 / 주요 기술 / 개발 내용 / 결과 및 분석)을
제목·소제목·불릿 구조로 정리한다.
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding="utf-8")

OUT = "docs/RansomGuard_캡스톤디자인_결과_내용.docx"

FONT = "맑은 고딕"
NAVY = RGBColor(0x16, 0x26, 0x3B)
TEAL = RGBColor(0x0E, 0x76, 0x6E)
GRAY = RGBColor(0x44, 0x4A, 0x55)

doc = Document()

# 기본 글꼴
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(11)
style.element.rPr.rFonts.set(
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)

# 여백 살짝 축소
for sec in doc.sections:
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)


def _set_font(run, size, color=NAVY, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)


def doc_title(text, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p.add_run(text), 22, NAVY, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(p2.add_run(sub), 12, GRAY)
    doc.add_paragraph()


def section(no, title):
    p = doc.add_paragraph()
    p.space_before = Pt(10)
    _set_font(p.add_run(f"{no}. "), 15, TEAL, bold=True)
    _set_font(p.add_run(title), 15, TEAL, bold=True)
    # 밑줄 역할의 얇은 단락
    bar = doc.add_paragraph()
    pPr = bar._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0E766E")
    pbdr.append(bottom)
    pPr.append(pbdr)
    bar.paragraph_format.space_after = Pt(6)


def sub(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    _set_font(p.add_run(text), 12, NAVY, bold=True)


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    _set_font(p.add_run(text), 11, GRAY)
    p.paragraph_format.space_after = Pt(2)


def plain(text):
    p = doc.add_paragraph()
    _set_font(p.add_run(text), 11, GRAY)


def _shade(cell, hex_color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(headers, rows, widths=None, caption=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        _set_font(p.add_run(caption), 10.5, GRAY, bold=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 헤더
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        _set_font(run, 10.5, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, "0E766E")
    # 본문
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(val)
            _set_font(run, 10, GRAY)
            if r % 2 == 1:
                _shade(cells[j], "EEF4F3")
            if j == 0:
                run.font.bold = True
                run.font.color.rgb = NAVY
    # 열 너비
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ════════════════════════════════════════════════════════════════════
doc_title("RansomGuard EDR",
          "행위 기반 실시간 랜섬웨어 탐지·차단 솔루션 (Windows 11) · 캡스톤디자인 결과 항목별 내용")

# ── 1. 개발 동기 및 목적 ───────────────────────────────────────────
section("1", "개발 동기 및 목적")
sub("■ 개발 동기")
bullet("랜섬웨어 피해가 매년 급증하고 있으나, 시그니처(이름표) 기반의 기존 백신은 "
       "이전에 본 적 없는 신종·변종·제로데이 공격을 탐지하지 못한다.")
bullet("파일이 일단 암호화되고 나면 복구가 사실상 불가능하므로, 사후 대응이 아니라 "
       "'암호화가 일어나는 순간'을 실시간으로 포착·차단하는 기술이 필요하다.")
bullet("기존 방식은 오탐(정상 작업을 악성으로 오인)과 탐지 지연 문제를 동시에 안고 있어, "
       "실사용 환경에 적용하기 어렵다.")
sub("■ 개발 목적")
bullet("알려지지 않은 신종 랜섬웨어까지 '무엇인가'가 아니라 '무슨 행동을 하는가(행위)'를 "
       "기준으로 탐지한다.")
bullet("사람의 개입 없이 위협을 자동으로 차단·격리하는 Windows 11 전용 "
       "EDR(엔드포인트 탐지·대응) 솔루션을 개발한다.")
bullet("단일 신호 하나로 단정하지 않고 여러 행위 단서를 종합 판단하여, "
       "탐지율은 높이고 오탐은 최소화하는 것을 목표로 한다.")
add_table(
    ["구분", "시그니처 기반 (기존)", "행위 기반 (본 과제)"],
    [
        ["탐지 기준", "알려진 악성코드 패턴(이름표)", "프로그램의 실제 동작·행위"],
        ["신종·제로데이", "탐지 불가", "탐지 가능"],
        ["대응 시점", "사후 (암호화 완료 후)", "실시간 (암호화 진행 중)"],
        ["복구 가능성", "사실상 불가", "암호화 차단으로 피해 최소화"],
    ],
    widths=[1.2, 2.8, 2.8],
    caption="[표 1] 기존 시그니처 방식과 행위 기반 방식 비교")

# ── 2. 주요 기술 ───────────────────────────────────────────────────
section("2", "주요 기술")
sub("■ 커널 미니필터 드라이버 (C, RansomGuard.sys)")
bullet("운영체제 커널 단계에서 모든 파일 입출력(I/O)을 실시간으로 후킹하여, "
       "유저 영역에서 우회하기 어려운 위치에서 행위를 수집한다.")
sub("■ 7종 행위 기반 탐지 엔진")
add_table(
    ["No.", "탐지 단서", "설명"],
    [
        ["1", "엔트로피 급증", "암호화된 파일의 무질서도(고엔트로피) 급증을 포착"],
        ["2", "미끼(디셉션) 파일 변조", "사용자가 건드리지 않는 가짜 파일 변경 시 즉시 악성 판정"],
        ["3", "대량 파일 변경", "단시간 내 다수 파일이 연속 수정되는 패턴 탐지"],
        ["4", "확장자 변조", "파일 확장자가 비정상적으로 일괄 변경되는 행위 탐지"],
        ["5", "백업·볼륨 섀도 삭제", "복원 무력화(VSS 삭제) 시도 포착"],
        ["6", "랜섬노트 생성", "금전 요구용 안내문 파일 생성 행위 탐지"],
        ["7", "위험 명령 실행", "권한 상승·복구 무력화 등 위험 명령 실행 탐지"],
    ],
    widths=[0.5, 1.9, 4.4],
    caption="[표 2] 7종 행위 기반 탐지 단서")
sub("■ 점수 합산 판단 엔진 (Python)")
bullet("2분 슬라이딩 윈도우 안에서 단서별 가중치를 합산하여 5단계 위험 등급을 판정한다.")
bullet("단일 단서로 단정하지 않는 '점수 합산' 방식으로 오탐을 최소화한다.")
sub("■ 자동 대응 및 자기 보호(Self-Protection)")
bullet("위험 등급 도달 시 악성 프로세스를 강제 종료·격리하고, 보안 솔루션 자신에 대한 "
       "변조를 방지하며, 사건 보고서를 자동 생성한다.")
sub("■ 실시간 대시보드 및 알림")
bullet("Flask 기반 한국어 실시간 대시보드로 탐지 현황을 시각화하고 사용자에게 알림을 제공한다.")

# ── 3. 개발 내용 ───────────────────────────────────────────────────
section("3", "개발 내용")
sub("■ 4단계 실시간 파이프라인 구현 (수집 → 분석 → 판단 → 대응)")
bullet("커널 드라이버 ↔ 유저 영역 탐지 엔진 간 통신 구조를 설계·구현했다.")
bullet("7종 탐지기를 모듈화하고 가중치 기반으로 통합하여 하나의 위험 점수로 합산했다.")
add_table(
    ["단계", "구성요소", "동작"],
    [
        ["① 수집", "커널 미니필터 (C)", "모든 파일 입출력(I/O)을 커널 단계에서 실시간 포착"],
        ["② 분석", "7종 행위 탐지기 (Python)", "엔트로피·미끼·대량변경 등 7종 단서 탐지"],
        ["③ 판단", "점수 합산 엔진", "2분 윈도우로 가중치 합산 → 5단계 위험 등급 판정"],
        ["④ 대응", "자동 대응 모듈", "프로세스 격리·종료 + 변조 방지 + 사건 보고서 생성"],
    ],
    widths=[0.9, 2.1, 3.8],
    caption="[표 3] 수집 → 분석 → 판단 → 대응 4단계 파이프라인")
sub("■ 탐지 정확도 튜닝")
bullet("엔트로피 임계치와 단서별 가중치를 반복 조정하여 오탐(FP)과 미탐(FN)의 균형을 맞췄다.")
sub("■ 검증 환경 구축")
bullet("네트워크 격리 VM과 안전 시뮬레이터로 대량 암호화·미끼 변조·백업 삭제 등 "
       "공격 시나리오를 재현했다.")
bullet("실제 랜섬웨어 검체(real-sample)를 이용한 실험 환경을 구성했다. "
       "(BSOD 발생 시에도 디스크 기록으로 결과 집계)")
sub("■ 운영 편의 기능")
bullet("한국어 실시간 대시보드, 알림 플러딩(과다 알림) 방지 로직을 구현했다.")
bullet("오탐·미탐(FP/FN) 자동 벤치마크 하니스와 성능 지표 PPT 자동 생성 도구를 개발했다.")

# ── 4. 결과 및 분석 ────────────────────────────────────────────────
section("4", "결과 및 분석")
sub("■ 수행 결과")
bullet("네트워크 격리 VM에서 실제 랜섬웨어 검체를 대상으로 탐지·차단에 성공했다.")
bullet("위험 등급 판정부터 자동 대응까지 2분 이내에 수행되었다.")
bullet("보호 프로세스 오종료 0건으로, 안전장치가 정상 작동함을 확인했다.")
sub("■ 정량 평가 (FP/FN 벤치마크)")
add_table(
    ["지표", "결과", "의미"],
    [
        ["오탐률 (FPR)", "0%", "정상 작업을 악성으로 오인한 사례 없음"],
        ["미탐률 (FNR)", "9.1%", "실제 악성 행위를 놓친 비율 (낮을수록 우수)"],
        ["F1-score", "95.2%", "정밀도·재현율의 조화 평균 (종합 성능)"],
        ["대응 시간", "2분 이내", "위험 등급 판정 → 자동 대응까지 소요 시간"],
        ["보호 프로세스 오종료", "0건", "안전장치가 정상 작동, 정상 프로세스 피해 없음"],
    ],
    widths=[1.8, 1.2, 3.8],
    caption="[표 4] 성능 평가 결과")
sub("■ 분석")
bullet("여러 행위 단서를 합산하는 점수 기반 판정이 단일 신호 방식 대비 오탐을 크게 "
       "줄인다는 점을 정량적으로 확인했다.")
bullet("커널 단계에서 행위를 수집함으로써 유저 영역 우회에 강인한 탐지 구조를 확보했다.")
sub("■ 한계 및 개선 방향")
bullet("정상 프로그램 화이트리스트를 강화하여 오탐을 추가로 감소시킬 필요가 있다.")
bullet("수 시간에 걸쳐 천천히 진행되는 '느린 암호화' 공격까지 포착하도록 확장이 필요하다.")
bullet("유저모드 PID 귀속(어느 프로세스가 원인인지 식별) 정밀도를 개선할 여지가 있다.")

doc.save(OUT)
print("saved:", OUT)
