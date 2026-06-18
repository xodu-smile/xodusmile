# -*- coding: utf-8 -*-
"""과제결과보고 최종발표 PPT 생성기 (RansomGuard EDR).

워드 파일 「과제결과보고 목차 및 콘텐츠(안).docx」의 목차/구성을 그대로 따른다.
  ✅ 서론  : 1.서론(1-1 배경/1-2 한계/1-3 목표) · 2.이론적 배경 · 3.진행 및 일정(3-1 R&R/3-2 일정)
  ✅ 본론  : 4.상세 아키텍처 · 5.과제 수행(5-1~5-3) · 6.과제 실험(6-1~6-3)
  ✅ 결론  : 7.결과(7-1~7-3) · 8.참고문헌
마지막은 시연영상 첨부 양식.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- 색상 테마 (라이트 / 고가시성) ----
BG    = RGBColor(0xF4, 0xF7, 0xFB)   # 슬라이드 배경 (밝은 회색)
CARD  = RGBColor(0xFF, 0xFF, 0xFF)   # 카드 (흰색)
LINE  = RGBColor(0xD8, 0xDF, 0xEA)   # 카드 테두리
INK   = RGBColor(0x16, 0x26, 0x3B)   # 본문/제목 (진한 잉크색)
MUTED = RGBColor(0x52, 0x62, 0x78)   # 보조 텍스트 (중간 회색)
NAVY  = RGBColor(0x16, 0x26, 0x3B)   # 강조색 위 텍스트 / 도형용 진한색
RED   = RGBColor(0xD3, 0x2F, 0x2F)
TEAL  = RGBColor(0x0E, 0x93, 0x86)
AMBER = RGBColor(0xB9, 0x73, 0x00)
BLUE  = RGBColor(0x25, 0x63, 0xC9)
# 하위 호환용 별칭 (본문에서 사용)
WHITE = INK
GRAY  = MUTED
NAVY2 = CARD

FONT = "맑은 고딕"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def txt(s, x, y, w, h, text, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, font=FONT, line_spacing=1.0, italic=False):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        f = r.font
        f.size = Pt(size); f.bold = bold; f.italic = italic
        f.color.rgb = color; f.name = font
    return tb


def card(s, x, y, w, h, fill=CARD, line=None):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.fill.solid(); c.fill.fore_color.rgb = fill
    # 라이트 배경에서 흰 카드가 묻히지 않도록 기본 테두리를 항상 둔다.
    c.line.color.rgb = line if line else LINE
    c.line.width = Pt(1.5 if line else 1.0)
    c.shadow.inherit = False
    return c


def bar(s, x, y, w, h, color):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    b.fill.solid(); b.fill.fore_color.rgb = color
    b.line.fill.background(); b.shadow.inherit = False
    return b


def header(s, kicker, title, accent=TEAL):
    bar(s, Inches(0.6), Inches(0.55), Inches(0.16), Inches(0.95), accent)
    txt(s, Inches(0.95), Inches(0.5), Inches(11.5), Inches(0.4), kicker, 14, accent, bold=True)
    txt(s, Inches(0.95), Inches(0.85), Inches(11.8), Inches(0.8), title, 30, WHITE, bold=True)


def divider(part_no, part_kr, subtitle, accent):
    """대분류(서론/본론/결론) 구분 슬라이드."""
    s = slide()
    bar(s, 0, Inches(3.0), SW, Inches(0.06), accent)
    txt(s, Inches(1), Inches(2.25), Inches(11.3), Inches(0.6), part_no, 20, accent, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1), Inches(2.75), Inches(11.3), Inches(1.1), part_kr, 52, WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1), Inches(4.2), Inches(11.3), Inches(0.6), subtitle, 17, GRAY, align=PP_ALIGN.CENTER)
    return s


def three_cards(s, items, y0=Inches(2.5), ch=Inches(3.1)):
    """3-열 카드 (번호·제목·본문)."""
    cw = Inches(3.75); gap = Inches(0.3); x0 = Inches(0.95)
    for i, (no, t, d, col) in enumerate(items):
        x = x0 + i * (cw + gap)
        card(s, x, y0, cw, ch, NAVY2, line=col)
        bar(s, x, y0, cw, Inches(0.12), col)
        txt(s, x + Inches(0.3), y0 + Inches(0.35), cw - Inches(0.6), Inches(0.4), no, 14, col, bold=True)
        txt(s, x + Inches(0.3), y0 + Inches(0.8), cw - Inches(0.6), Inches(0.6), t, 19, WHITE, bold=True)
        txt(s, x + Inches(0.3), y0 + Inches(1.55), cw - Inches(0.6), ch - Inches(1.7), d, 14, GRAY, line_spacing=1.25)


# =====================================================================
# 표지
# =====================================================================
s = slide()
bar(s, 0, Inches(3.05), SW, Inches(0.05), TEAL)
txt(s, Inches(1), Inches(1.55), Inches(11.3), Inches(0.5), "과제결과보고 · 최종 발표", 18, TEAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(2.1), Inches(11.3), Inches(1.0), "RansomGuard EDR", 54, WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(3.2), Inches(11.3), Inches(0.7),
    "행위 기반 실시간 랜섬웨어 탐지·차단 솔루션", 22, GRAY, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(4.4), Inches(11.3), Inches(0.5),
    "Windows 11 전용 · 커널 행위 탐지 + 자동 대응", 15, BLUE, align=PP_ALIGN.CENTER)
txt(s, Inches(1), Inches(6.35), Inches(11.3), Inches(0.5),
    "팀명 / 발표자 · 지도교수 [성함] · 2026", 14, GRAY, align=PP_ALIGN.CENTER)

# =====================================================================
# 목차
# =====================================================================
s = slide()
header(s, "CONTENTS", "목차", BLUE)
cols = [
    ("✅  서론", RED, [
        "1. 서론  (배경 · 한계 · 목표)",
        "2. 이론적 배경",
        "3. 과제 진행 및 일정  (R&R · 일정)",
    ]),
    ("✅  본론", AMBER, [
        "4. 상세 아키텍처",
        "5. 과제 수행  (구현 · 과정 · 실행)",
        "6. 과제 실험  (환경 · 악성 · 정상)",
    ]),
    ("✅  결론", TEAL, [
        "7. 결과  (성과 · 한계 · 소감)",
        "8. 참고문헌",
        "▶  시연 영상",
    ]),
]
cw = Inches(3.75); gap = Inches(0.3); x0 = Inches(0.95); y0 = Inches(2.1)
for i, (title, col, rows) in enumerate(cols):
    x = x0 + i * (cw + gap)
    card(s, x, y0, cw, Inches(4.4), NAVY2)
    bar(s, x, y0, cw, Inches(0.12), col)
    txt(s, x + Inches(0.3), y0 + Inches(0.35), cw - Inches(0.6), Inches(0.5), title, 20, col, bold=True)
    for j, r in enumerate(rows):
        txt(s, x + Inches(0.3), y0 + Inches(1.2) + j * Inches(0.95), cw - Inches(0.55), Inches(0.9),
            r, 15, WHITE, line_spacing=1.1)

# =====================================================================
# ✅ 서론
# =====================================================================
divider("PART 1", "서론", "왜 이 과제가 필요한가 — 배경 · 한계 · 목표 · 이론 · 일정", RED)

# ---- 1. 서론 (1-1 / 1-2 / 1-3) ----
s = slide()
header(s, "1. 서론", "기존 탐지의 한계와 행위 기반 탐지의 필요성", RED)
txt(s, Inches(0.95), Inches(1.8), Inches(11.5), Inches(0.55),
    "랜섬웨어 위협이 급증하는 환경에서 시그니처 기반 탐지의 한계를 분석하고, 행위 기반 탐지의 필요성을 제시한다.",
    16, GRAY, line_spacing=1.2)
three_cards(s, [
    ("1-1  프로젝트 배경", "시그니처의 한계",
     "시그니처(이름표) 기반 탐지는 제로데이·신종에 무력하다. 실시간 '행위' 탐지 기술의 필요성이 커지고 있다.", BLUE),
    ("1-2  기존 한계점", "오탐·탐지 지연",
     "기존 방식은 오탐과 탐지 지연 문제를 안고 있으며, 특히 변형·신종 랜섬웨어 대응에 취약하다.", AMBER),
    ("1-3  과제 목표", "탐지 + 즉시 차단",
     "실시간 행위 분석으로 신종 랜섬웨어를 탐지·차단하는 솔루션 개발. 범위와 한계를 명확히 설정한다.", TEAL),
], y0=Inches(2.5), ch=Inches(3.4))

# ---- 2. 이론적 배경 ----
s = slide()
header(s, "2. 이론적 배경", "랜섬웨어 동작 원리와 핵심 탐지 개념", BLUE)
rows = [
    ("🔓", "랜섬웨어 동작 원리", "초기 침투 → 권한 상승 → 백업·복원 무력화 → 대량 파일 암호화 → 금전 요구의 단계로 전개된다."),
    ("👁", "행위 기반 탐지", "'무엇인가(시그니처)'가 아니라 '무슨 행동을 하는가'를 본다. 미지의 신종도 행동으로 포착할 수 있다."),
    ("📊", "엔트로피 분석", "정상 파일은 규칙적, 암호화된 파일은 무질서(고엔트로피)하다. 엔트로피 급증은 대량 암호화의 핵심 신호다."),
    ("🪤", "디셉션(미끼 파일)", "사용자가 건드리지 않는 가짜 파일을 배치한다. 이 파일이 변경되면 악성 행위로 즉시 판정한다."),
]
y = Inches(2.0)
for i, (ic, t, d) in enumerate(rows):
    yy = y + i * Inches(1.15)
    card(s, Inches(0.95), yy, Inches(11.45), Inches(0.98), NAVY2)
    txt(s, Inches(1.2), yy, Inches(0.95), Inches(0.98), ic, 28, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    txt(s, Inches(2.2), yy + Inches(0.13), Inches(3.4), Inches(0.7), t, 17, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(5.7), yy, Inches(6.5), Inches(0.98), d, 14, GRAY, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)

# ---- 3. 과제 진행 및 일정 (3-1 R&R / 3-2 추진 일정) ----
s = slide()
header(s, "3. 과제 진행 및 일정", "역할 분담(R&R)과 추진 일정", AMBER)

# 3-1 R&R (좌측)
txt(s, Inches(0.95), Inches(1.8), Inches(5.6), Inches(0.4), "3-1  R&R (역할 분담)", 17, TEAL, bold=True)
rr = [
    ("분석 · 설계", "위협 분석, 탐지 로직·아키텍처 설계", BLUE),
    ("구현", "커널 드라이버 · 탐지 엔진 · 자동 대응 개발", AMBER),
    ("성능 테스트", "VM 실험, 탐지·오탐 검증, 보고서 작성", TEAL),
]
for i, (t, d, col) in enumerate(rr):
    y = Inches(2.35) + i * Inches(1.05)
    card(s, Inches(0.95), y, Inches(5.6), Inches(0.9), NAVY2)
    bar(s, Inches(0.95), y, Inches(0.12), Inches(0.9), col)
    txt(s, Inches(1.3), y + Inches(0.12), Inches(5.0), Inches(0.4), t, 16, WHITE, bold=True)
    txt(s, Inches(1.3), y + Inches(0.5), Inches(5.1), Inches(0.35), d + "  · [팀원 성함]", 12.5, GRAY)

# 3-2 추진 일정 (우측 타임라인)
txt(s, Inches(6.75), Inches(1.8), Inches(5.6), Inches(0.4), "3-2  추진 일정 (6주)", 17, AMBER, bold=True)
plan = [
    ("1–2주", "설계 · 환경 구축", BLUE),
    ("3–4주", "개발 (드라이버 · 엔진 · 대응)", AMBER),
    ("5주", "통합 · 시뮬레이터 검증", TEAL),
    ("6주", "실검체 실험 · 보고서", RED),
]
for i, (wk, d, col) in enumerate(plan):
    y = Inches(2.35) + i * Inches(1.0)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.95), y + Inches(0.12), Inches(0.3), Inches(0.3))
    dot.fill.solid(); dot.fill.fore_color.rgb = col; dot.line.fill.background(); dot.shadow.inherit = False
    if i < 3:
        bar(s, Inches(7.08), y + Inches(0.42), Inches(0.04), Inches(0.7), GRAY)
    txt(s, Inches(7.45), y, Inches(1.3), Inches(0.55), wk, 15, col, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(8.75), y, Inches(3.6), Inches(0.55), d, 14, GRAY, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# ✅ 본론
# =====================================================================
divider("PART 2", "본론", "어떻게 만들었나 — 아키텍처 · 수행 · 실험", AMBER)

# ---- 4. 상세 아키텍처 ----
s = slide()
header(s, "4. 상세 아키텍처", "수집 → 분석 → 판단 → 대응의 4단계 구조", BLUE)
txt(s, Inches(0.95), Inches(1.8), Inches(11.5), Inches(0.5),
    "커널(시스템 깊은 곳)에서 행위를 수집하고, 유저 영역의 엔진이 점수를 합산해 위험 등급을 판정·대응한다.", 16, GRAY)
steps = [
    ("1", "수집", "커널 미니필터가\n모든 파일 변경을 포착", BLUE),
    ("2", "분석", "엔트로피·미끼·위험명령\n등 7종 단서 탐지", AMBER),
    ("3", "판단", "2분 창에서 점수 합산\n→ 5단계 위험 등급", RED),
    ("4", "대응", "격리·종료 + 변조방지\n+ 사건 보고서 작성", TEAL),
]
cw = Inches(2.75); gap = Inches(0.35); x0 = Inches(0.7); y0 = Inches(2.75)
for i, (n, t, d, col) in enumerate(steps):
    x = x0 + i * (cw + gap)
    card(s, x, y0, cw, Inches(2.5), NAVY2)
    bar(s, x, y0, cw, Inches(0.12), col)
    cnum = s.shapes.add_shape(MSO_SHAPE.OVAL, x + cw/2 - Inches(0.4), y0 + Inches(0.35), Inches(0.8), Inches(0.8))
    cnum.fill.solid(); cnum.fill.fore_color.rgb = col; cnum.line.fill.background(); cnum.shadow.inherit = False
    p = cnum.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = n; r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = FONT
    txt(s, x, y0 + Inches(1.35), cw, Inches(0.5), t, 21, WHITE, bold=True, align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.2), y0 + Inches(1.9), cw - Inches(0.4), Inches(0.55), d, 13, GRAY, align=PP_ALIGN.CENTER, line_spacing=1.1)
    if i < 3:
        txt(s, x + cw - Inches(0.05), y0 + Inches(0.9), Inches(0.5), Inches(0.6), "➜", 26, GRAY, align=PP_ALIGN.CENTER)
txt(s, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.9),
    "커널 드라이버(C, RansomGuard.sys) ↔ 탐지 엔진(Python) ↔ 대시보드/보고서  ·  단서 하나로 단정하지 않고 점수 합산으로 오탐을 줄인다.",
    15, TEAL, bold=True, align=PP_ALIGN.CENTER, line_spacing=1.2)

# ---- 5. 과제 수행 (5-1 / 5-2 / 5-3) ----
s = slide()
header(s, "5. 과제 수행", "솔루션 구현 · 구현 과정 · 솔루션 실행", TEAL)
three_cards(s, [
    ("5-1  솔루션 구현", "수집–분석–판단–대응",
     "커널 미니필터 + 7종 행위 탐지기 + 점수 채점 엔진 + 자동 대응(격리·종료)으로 구성된 아키텍처를 구현했다.", BLUE),
    ("5-2  구현 과정", "알고리즘 · 임계치 조정",
     "엔트로피·미끼·대량변경 등 탐지 알고리즘을 설계하고, 오탐과 미탐의 균형을 맞추도록 임계치를 반복 조정했다.", AMBER),
    ("5-3  솔루션 실행", "알림 · 로그 · 보고서",
     "탐지 결과를 사용자 알림과 로그로 남기고, 실시간 대시보드와 사건 보고서를 자동 생성하도록 구현했다.", TEAL),
], y0=Inches(2.4), ch=Inches(3.5))

# ---- 6. 과제 실험 (6-1 / 6-2 / 6-3) ----
s = slide()
header(s, "6. 과제 실험", "환경 구축 · 악성 행위 탐지 · 정상 행위 미탐지", AMBER)
three_cards(s, [
    ("6-1  과제 환경", "격리 VM",
     "네트워크 격리 VM으로 실제 환경과 유사한 조건을 구성했다. BSOD가 나도 디스크 기록으로 결과를 집계한다.", BLUE),
    ("6-2  랜섬웨어 실험", "악성 행위 탐지",
     "안전 시뮬레이터와 실검체로 대량 암호화·미끼 변조·백업 삭제 등 시나리오의 탐지·차단 성능을 검증했다.", RED),
    ("6-3  정상행위 실험", "오탐 검증",
     "압축·대량 파일 변경 등 정상 작업으로 오탐 여부를 검증하고, 보호 프로세스 오종료가 없음을 확인했다.", TEAL),
], y0=Inches(2.4), ch=Inches(3.5))

# =====================================================================
# ✅ 결론
# =====================================================================
divider("PART 3", "결론", "무엇을 이뤘나 — 결과 · 한계 · 소감 · 참고문헌", TEAL)

# ---- 7. 결과 (7-1 / 7-2 / 7-3) ----
s = slide()
header(s, "7. 결과", "수행 결과 · 한계점 및 개선점 · 소감", TEAL)

# 7-1 성과 지표
txt(s, Inches(0.95), Inches(1.75), Inches(11.5), Inches(0.4), "7-1  과제 수행 결과", 17, TEAL, bold=True)
stats = [
    ("실검체 탐지", "격리 VM에서 실제\n랜섬웨어 탐지 성공", TEAL),
    ("2분 이내", "위험 등급 판정 →\n자동 대응까지", BLUE),
    ("오종료 0건", "보호 프로세스\n안전장치 정상 작동", AMBER),
]
cw = Inches(3.75); x0 = Inches(0.95); y0 = Inches(2.25)
for i, (big, d, col) in enumerate(stats):
    x = x0 + i * (cw + Inches(0.3))
    card(s, x, y0, cw, Inches(1.65), NAVY2)
    bar(s, x, y0, cw, Inches(0.1), col)
    txt(s, x, y0 + Inches(0.3), cw, Inches(0.6), big, 26, col, bold=True, align=PP_ALIGN.CENTER)
    txt(s, x + Inches(0.2), y0 + Inches(1.0), cw - Inches(0.4), Inches(0.55), d, 13, GRAY, align=PP_ALIGN.CENTER, line_spacing=1.05)

# 7-2 / 7-3
card(s, Inches(0.95), Inches(4.2), Inches(5.6), Inches(2.55), NAVY2, line=AMBER)
txt(s, Inches(1.2), Inches(4.4), Inches(5.1), Inches(0.4), "7-2  한계점 및 개선점", 16, AMBER, bold=True)
txt(s, Inches(1.2), Inches(4.9), Inches(5.15), Inches(1.8),
    "•  정상 프로그램 화이트리스트 강화로 오탐 추가 감소\n"
    "•  수 시간에 걸친 '느린 암호화'까지 포착\n"
    "•  포렌식을 위한 커널 단계 컨텍스트 저장 확대\n"
    "•  유저모드 PID 귀속 정밀도 개선",
    13.5, GRAY, line_spacing=1.3)
card(s, Inches(6.75), Inches(4.2), Inches(5.6), Inches(2.55), NAVY2, line=TEAL)
txt(s, Inches(7.0), Inches(4.4), Inches(5.1), Inches(0.4), "7-3  소감", 16, TEAL, bold=True)
txt(s, Inches(7.0), Inches(4.9), Inches(5.15), Inches(1.8),
    "커널 영역의 행위 탐지를 직접 구현하며 OS 내부와 보안 동작 원리를 깊이 이해했다. "
    "탐지·대응·안전장치를 통합하는 과정에서 역할 분담과 협업의 중요성을 체감했다.",
    13.5, GRAY, line_spacing=1.3)

# ---- 8. 참고문헌 ----
s = slide()
header(s, "8. 참고문헌", "과제 수행 근거 자료", BLUE)
refs = [
    "MITRE ATT&CK — Data Encrypted for Impact (T1486), Inhibit System Recovery (T1490)",
    "Microsoft Docs — File System Minifilter Drivers / Filter Manager Concepts",
    "Microsoft Docs — Windows Driver Kit (WDK) 및 커널 모드 드라이버 아키텍처",
    "C. E. Shannon, \"A Mathematical Theory of Communication\" (엔트로피 이론)",
    "Kharraz et al., \"Cutting the Gordian Knot: A Look Under the Hood of Ransomware Attacks\"",
    "NIST SP 800-83 — Guide to Malware Incident Prevention and Handling",
]
for i, r in enumerate(refs):
    y = Inches(1.95) + i * Inches(0.82)
    card(s, Inches(0.95), y, Inches(11.45), Inches(0.68), NAVY2)
    txt(s, Inches(1.2), y, Inches(0.6), Inches(0.68), f"[{i+1}]", 14, TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(1.85), y, Inches(10.4), Inches(0.68), r, 13.5, GRAY, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
txt(s, Inches(0.95), Inches(6.95), Inches(11.5), Inches(0.4),
    "※ 실제 인용한 자료로 교체하여 최종본을 완성하세요.", 12, GRAY, italic=True)

# =====================================================================
# 시연 영상 (placeholder 양식)
# =====================================================================
s = slide()
bar(s, 0, Inches(0.55), SW, Inches(0.05), TEAL)
txt(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.5), "DEMO", 16, TEAL, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.6), Inches(1.05), Inches(12), Inches(0.7), "시연 영상", 34, WHITE, bold=True, align=PP_ALIGN.CENTER)
vx, vy, vw, vh = Inches(2.4), Inches(2.1), Inches(8.5), Inches(4.0)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, vx, vy, vw, vh)
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x29, 0x3A)
box.line.color.rgb = TEAL; box.line.width = Pt(2); box.shadow.inherit = False
play = s.shapes.add_shape(MSO_SHAPE.OVAL, vx + vw/2 - Inches(0.7), vy + vh/2 - Inches(0.7), Inches(1.4), Inches(1.4))
play.fill.solid(); play.fill.fore_color.rgb = TEAL; play.line.fill.background(); play.shadow.inherit = False
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, vx + vw/2 - Inches(0.25), vy + vh/2 - Inches(0.35), Inches(0.6), Inches(0.7))
tri.rotation = 90
tri.fill.solid(); tri.fill.fore_color.rgb = NAVY; tri.line.fill.background(); tri.shadow.inherit = False
txt(s, vx, vy + vh - Inches(0.85), vw, Inches(0.6),
    "▶ 여기에 시연 영상을 삽입하세요  (삽입 ▸ 비디오 ▸ 이 디바이스의 비디오)",
    14, RGBColor(0xC9, 0xD3, 0xDE), align=PP_ALIGN.CENTER)
txt(s, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.9),
    "추천 시연 흐름:  ① 에이전트 실행 → ② 시뮬레이터로 대량 암호화 발생 → "
    "③ 대시보드 점수 급상승(CRITICAL) → ④ 자동 종료 + 사건 보고서 생성",
    14, GRAY, align=PP_ALIGN.CENTER, line_spacing=1.2)

out = "docs/RansomGuard_과제결과보고_최종발표.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides._sldIdLst))
