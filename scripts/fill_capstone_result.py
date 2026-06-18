# -*- coding: utf-8 -*-
"""캡스톤디자인 결과 템플릿(포스터)에 RansomGuard EDR 내용을 채운다.

레이아웃·서체(KoPub돋움체 Light)·색상은 그대로 두고, 각 항목 본문
텍스트박스의 placeholder 만 프로젝트 내용으로 교체한다.

  · 개발 동기 및 목적  → 본문 id=40
  · 주요 기술          → 본문 id=52 (중복 잔여 박스 id=37 은 비움)
  · 개발 내용          → 본문 id=47
  · 결과 및 분석       → 본문 id=44
  · 과제명(제목)       → id=4

주의: 템플릿의 우/하단 본문 박스는 wrap="none"(줄바꿈 안 함) + 폭 6in 로
저작돼 있어 긴 본문이 박스 밖으로 흘러 보이지 않는다. 따라서 네 본문 박스
모두 wrap="square"(자동 줄바꿈) + 동일 폭/폰트로 통일한다.
"""
import sys
from copy import deepcopy
import lxml.etree as etree
from pptx import Presentation
from pptx.util import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

SRC = "2026학년도 1학기 캡스톤디자인 결과 탬플릿.pptx"
OUT = "2026학년도 1학기 캡스톤디자인 결과 탬플릿.pptx"

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

prs = Presentation(SRC)
slide = prs.slides[0]


def find(shapes, target_id):
    for sh in shapes:
        if sh.shape_id == target_id:
            return sh
        if sh.shape_type == 6:  # group
            r = find(sh.shapes, target_id)
            if r is not None:
                return r
    return None


def _rewrite(tf, lines, sz):
    """텍스트프레임 본문을 lines 로 교체(한 문단 + <a:br> 줄바꿈).
    첫 run 의 서식을 복제하되 글자 크기(sz, 1/100pt)만 통일한다."""
    base_r = tf.paragraphs[0].runs[0]
    rPr = base_r._r.find(A + 'rPr')
    txBody = tf._txBody
    for p in txBody.findall(A + 'p')[1:]:
        txBody.remove(p)
    first_p = txBody.findall(A + 'p')[0]
    for child in list(first_p):
        if child.tag in (A + 'r', A + 'br'):
            first_p.remove(child)

    def make_run(text):
        r = etree.SubElement(first_p, A + 'r')
        if rPr is not None:
            rp = deepcopy(rPr)
            if sz is not None:
                rp.set('sz', str(sz))
            r.append(rp)
        t = etree.SubElement(r, A + 't')
        t.text = text

    for i, ln in enumerate(lines):
        if i > 0:
            etree.SubElement(first_p, A + 'br')
        make_run(ln)


def set_body(shape_id, lines, width_in=13.4, height_in=14.0, sz=3000):
    sh = find(slide.shapes, shape_id)
    _rewrite(sh.text_frame, lines, sz)
    # 자동 줄바꿈 켜기
    bodyPr = sh.text_frame._txBody.find(A + 'bodyPr')
    bodyPr.set('wrap', 'square')
    # 폭/높이 통일 (그룹 scaleX=1 이라 자식 ext = 슬라이드 폭)
    if width_in:
        sh.width = Inches(width_in)
    if height_in:
        sh.height = Inches(height_in)


def set_title(shape_id, line1, line2):
    sh = find(slide.shapes, shape_id)
    _rewrite(sh.text_frame, [line1, line2] if line2 else [line1], None)


# ── 과제명 ────────────────────────────────────────────────────────────
set_title(4, "RansomGuard EDR",
          "행위 기반 실시간 랜섬웨어 탐지·차단 솔루션 (Windows 11)")

# ── 개발 동기 및 목적 ─────────────────────────────────────────────────
set_body(40, [
    "• 랜섬웨어 피해가 매년 급증하지만, 시그니처(이름표) 기반 백신은",
    "  신종·변종·제로데이 공격을 잡지 못한다.",
    "• 파일이 암호화된 뒤에는 복구가 사실상 불가능하므로,",
    "  '암호화가 일어나는 순간'을 실시간으로 포착·차단해야 한다.",
    "",
    "[목적]",
    "• 알려지지 않은 신종까지 '무슨 행동을 하는가(행위)'로 탐지하고,",
    "  사람 개입 없이 자동으로 차단·격리하는 Windows 11 전용 EDR",
    "  (엔드포인트 탐지·대응) 솔루션을 개발한다.",
    "• 단일 신호로 단정하지 않고 여러 행위 단서를 종합 판단하여,",
    "  탐지율은 높이고 오탐(정상 작업 차단)은 최소화한다.",
])

# ── 주요 기술 ─────────────────────────────────────────────────────────
set_body(52, [
    "• 커널 미니필터 드라이버 (C, RansomGuard.sys)",
    "  – 모든 파일 입출력(I/O)을 커널 단계에서 실시간 후킹",
    "• 7종 행위 기반 탐지 엔진",
    "  – 엔트로피 급증, 미끼(디셉션) 파일 변조, 대량 파일 변경,",
    "    확장자 변조, 백업·볼륨 섀도 삭제, 랜섬노트 생성, 위험 명령",
    "• 점수 합산 판단 엔진 (Python)",
    "  – 2분 슬라이딩 윈도우로 단서별 가중치 합산",
    "  – 5단계 위험 등급 판정 (단일 단서 단정 X → 오탐 최소화)",
    "• 자동 대응 + 자기 보호(self-protection)",
    "  – 악성 프로세스 강제 종료·격리, 변조 방지, 사건 보고서 자동 생성",
    "• 실시간 한국어 대시보드(Flask) · 사용자 알림",
])
# 우측 상단의 중복 잔여 placeholder 박스는 비운다.
set_body(37, [""], width_in=None, height_in=None, sz=None)

# ── 개발 내용 ─────────────────────────────────────────────────────────
set_body(47, [
    "• 수집 → 분석 → 판단 → 대응 4단계 실시간 파이프라인 구현",
    "  – 커널 드라이버 ↔ 유저 영역 탐지 엔진 간 통신 구조 설계",
    "  – 7종 탐지기를 모듈화하여 가중치 기반으로 통합",
    "• 엔트로피·임계치 반복 튜닝으로 오탐·미탐의 균형 조정",
    "• 검증 환경 구축",
    "  – 네트워크 격리 VM + 안전 시뮬레이터로 공격 시나리오 재현",
    "  – 실제 랜섬웨어 검체(real-sample) 실험 환경 구성",
    "• 운영 편의 기능",
    "  – 한국어 실시간 대시보드, 알림 플러딩 방지 로직",
    "  – 오탐·미탐(FP/FN) 자동 벤치마크 및 성능 지표 PPT 생성",
])

# ── 결과 및 분석 ──────────────────────────────────────────────────────
set_body(44, [
    "• 격리 VM 실검체 테스트에서 실제 랜섬웨어 탐지·차단 성공",
    "",
    "[FP/FN 벤치마크 결과]",
    "  · 오탐률(FPR) 0%   · 미탐률(FNR) 9.1%   · F1-score 95.2%",
    "",
    "• 위험 등급 판정 → 자동 대응까지 2분 이내 수행",
    "• 보호 프로세스 오종료 0건 (안전장치 정상 작동)",
    "• 점수 합산 방식이 단일 신호 대비 오탐을 크게 줄임을 확인",
    "",
    "[한계 및 개선 방향]",
    "• 정상 프로그램 화이트리스트 강화로 오탐 추가 감소",
    "• 수 시간에 걸친 '느린 암호화'까지 포착",
    "• 유저모드 PID 귀속 정밀도 개선",
])

prs.save(OUT)
print("saved:", OUT)
